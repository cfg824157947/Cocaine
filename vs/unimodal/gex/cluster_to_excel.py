#!/usr/bin/env python3

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


# -------------------------
# Formatting helper
# -------------------------
def format_worksheet(ws):
    header_fill = PatternFill(fill_type="solid", start_color="000000", end_color="000000")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    default_font = Font(name="Arial")
    italic_font = Font(name="Arial", italic=True)

    # Header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    # Body
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column == 1:
                cell.font = italic_font
            else:
                cell.font = default_font

            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000"


# -------------------------
# Excel writer
# -------------------------
def write_excel(input_dir, clusters, pattern, outname):
    with pd.ExcelWriter(outname, engine="openpyxl") as writer:
        for x in clusters:
            infile = input_dir / pattern.format(x=x)
            if not infile.exists():
                continue

            df = pd.read_csv(infile, sep="\t")
            sheet_name = f"Cluster {x}"
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    # Apply formatting
    wb = load_workbook(outname)
    for ws in wb.worksheets:
        format_worksheet(ws)
    wb.save(outname)


# -------------------------
# Main
# -------------------------
def main(input_dir, output_prefix):
    input_dir = Path(input_dir)
    clusters = range(46)  # explicit: 0–45

    write_excel(
        input_dir=input_dir,
        clusters=clusters,
        pattern="cluster_{x}_lsmeans_Treatment_Line_Sex.txt",
        outname=f"{output_prefix}_TxLxS.xlsx",
    )

    write_excel(
        input_dir=input_dir,
        clusters=clusters,
        pattern="cluster_{x}_lsmeans_Treatment_Line.txt",
        outname=f"{output_prefix}_TxL.xlsx",
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Convert cluster lsmeans txt files into formatted Excel workbooks"
    )
    parser.add_argument(
        "-i", "--input_dir", required=True,
        help="Directory containing cluster lsmeans txt files"
    )
    parser.add_argument(
        "-o", "--output_prefix", required=True,
        help="Prefix for output Excel files"
    )

    args = parser.parse_args()
    main(args.input_dir, args.output_prefix)